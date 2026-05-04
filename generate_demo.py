import pandas as pd
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

demo_data = [
    {"name": "ივანე ჯავახიშვილის სახელობის თბილისის სახელმწიფო უნივერსიტეტი", "address": "ი. ჭავჭავაძის გამზ. 1", "phone": "032 2 22 51 07", "website": "https://www.tsu.ge", "google_maps_url": "https://goo.gl/maps/tsu", "lat": 41.7093, "lng": 44.7816, "rating": 4.5, "category": "University", "opening_hours": "09:00 - 18:00"},
    {"name": "საქართველოს ტექნიკური უნივერსიტეტი", "address": "მ. კოსტავას ქ. 77", "phone": "032 2 36 51 52", "website": "https://gtu.ge", "google_maps_url": "https://goo.gl/maps/gtu", "lat": 41.7165, "lng": 44.7785, "rating": 4.2, "category": "University", "opening_hours": "09:00 - 18:00"},
    {"name": "თავისუფალი უნივერსიტეტი", "address": "დ. აღმაშენებლის ხეივანი 240", "phone": "032 2 20 09 01", "website": "https://freeuni.edu.ge", "google_maps_url": "https://goo.gl/maps/freeuni", "lat": 41.8055, "lng": 44.7674, "rating": 4.9, "category": "University", "opening_hours": "10:00 - 19:00"},
    {"name": "53-ე საჯარო სკოლა", "address": "ბარნოვის ქ. 54", "phone": "032 2 99 90 12", "website": "N/A", "google_maps_url": "https://goo.gl/maps/school53", "lat": 41.7065, "lng": 44.7895, "rating": 4.0, "category": "School", "opening_hours": "08:30 - 16:00"},
    {"name": "ბაქსვუდის საერთაშორისო სკოლა", "address": "წყნეთი, რუსთაველის ქ. 32", "phone": "032 2 22 55 55", "website": "https://buckwood.ge", "google_maps_url": "https://goo.gl/maps/buckwood", "lat": 41.7025, "lng": 44.6955, "rating": 4.7, "category": "School", "opening_hours": "09:00 - 17:00"},
    {"name": "გივი ზალდასტანიშვილის სახელობის ამერიკული აკადემია", "address": "ლისის ტბის მიმდებარედ", "phone": "032 2 22 74 41", "website": "https://gzaat.ge", "google_maps_url": "https://goo.gl/maps/gzaat", "lat": 41.7455, "lng": 44.7355, "rating": 4.8, "category": "Academy/Learning Center", "opening_hours": "08:30 - 16:30"},
    {"name": "ბრიტანული საბავშვო ბაღი", "address": "თ. აბულაძის ქ. 12", "phone": "032 2 25 11 22", "website": "https://britishkindergarten.ge", "google_maps_url": "https://goo.gl/maps/bk", "lat": 41.7125, "lng": 44.7555, "rating": 4.6, "category": "Kindergarten", "opening_hours": "09:00 - 19:00"},
    {"name": "კავკასიის უნივერსიტეტი", "address": "პაატა სააკაძის ქ. 1", "phone": "032 2 37 77 77", "website": "https://cu.edu.ge", "google_maps_url": "https://goo.gl/maps/cu", "lat": 41.7255, "lng": 44.8155, "rating": 4.6, "category": "University", "opening_hours": "09:00 - 19:00"},
    {"name": "საქართველოს უნივერსიტეტი (UG)", "address": "მ. კოსტავას ქ. 77ა", "phone": "032 2 55 22 22", "website": "https://ug.edu.ge", "google_maps_url": "https://goo.gl/maps/ug", "lat": 41.7185, "lng": 44.7755, "rating": 4.4, "category": "University", "opening_hours": "09:00 - 20:00"}
]

def save_to_excel(data):
    df = pd.DataFrame(data)
    writer = pd.ExcelWriter('educational_institutions.xlsx', engine='openpyxl')
    
    # Summary Sheet
    summary_data = df['category'].value_counts().reset_index()
    summary_data.columns = ['Category', 'Count']
    summary_data.to_excel(writer, sheet_name='Summary', index=False)
    
    ws_summary = writer.sheets['Summary']
    for cell in ws_summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")

    # Category Sheets
    for cat in df['category'].unique():
        cat_df = df[df['category'] == cat]
        sheet_name = cat.replace('/', '-')[:30]
        cat_df.to_excel(writer, sheet_name=sheet_name, index=False)
        
        ws = writer.sheets[sheet_name]
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="2980B9", end_color="2980B9", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")

    writer.close()

def generate_html(data):
    html_template = """
<!DOCTYPE html>
<html lang="ka">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Tbilisi Education Explorer</title>
    <link rel="stylesheet" href="https://unpkg.com/leaflet@1.9.4/dist/leaflet.css" />
    <link href="https://fonts.googleapis.com/css2?family=Outfit:wght@300;400;600&display=swap" rel="stylesheet">
    <style>
        :root {
            --primary: #2563eb;
            --secondary: #1e293b;
            --bg: #f8fafc;
            --card-bg: rgba(255, 255, 255, 0.9);
            --accent: #3b82f6;
        }

        body { 
            font-family: 'Outfit', sans-serif; 
            margin: 0; 
            display: flex; 
            flex-direction: column; 
            height: 100vh; 
            background: var(--bg);
            color: var(--secondary);
        }

        #header { 
            background: var(--secondary); 
            color: white; 
            padding: 1rem 2rem; 
            display: flex; 
            justify-content: space-between; 
            align-items: center;
            box-shadow: 0 4px 6px -1px rgb(0 0 0 / 0.1);
            z-index: 1000;
        }

        #header h1 { margin: 0; font-size: 1.5rem; font-weight: 600; }

        #controls { 
            padding: 1rem 2rem; 
            background: white; 
            display: flex; 
            gap: 1rem; 
            align-items: center; 
            box-shadow: 0 1px 3px 0 rgb(0 0 0 / 0.1);
            z-index: 999;
        }

        #map { flex-grow: 1; z-index: 1; }

        .legend { 
            background: var(--card-bg); 
            padding: 1rem; 
            line-height: 2; 
            border-radius: 12px; 
            box-shadow: 0 10px 15px -3px rgb(0 0 0 / 0.1);
            backdrop-filter: blur(8px);
            border: 1px solid rgba(255,255,255,0.2);
        }

        .legend h4 { margin: 0 0 0.5rem 0; font-size: 0.9rem; text-transform: uppercase; letter-spacing: 0.05em; color: #64748b; }
        .legend i { 
            width: 12px; 
            height: 12px; 
            float: left; 
            margin-right: 10px; 
            margin-top: 6px;
            border-radius: 50%; 
            box-shadow: 0 0 0 2px #fff, 0 0 0 3px currentColor;
        }

        input, select, button { 
            padding: 0.6rem 1rem; 
            border-radius: 8px; 
            border: 1px solid #e2e8f0; 
            font-family: inherit;
            transition: all 0.2s;
        }

        input:focus, select:focus { outline: none; border-color: var(--primary); ring: 2px var(--primary); }

        button { 
            cursor: pointer; 
            background: var(--primary); 
            color: white; 
            border: none; 
            font-weight: 600;
            display: flex;
            align-items: center;
            gap: 0.5rem;
        }

        button:hover { background: #1d4ed8; transform: translateY(-1px); }

        .leaflet-popup-content-wrapper {
            border-radius: 12px;
            padding: 0.5rem;
        }

        .popup-card { min-width: 200px; }
        .popup-card h3 { margin: 0 0 0.5rem 0; font-size: 1.1rem; color: var(--secondary); }
        .popup-card p { margin: 0.25rem 0; font-size: 0.9rem; color: #64748b; }
        .popup-card a { color: var(--primary); text-decoration: none; font-weight: 500; }
        .popup-card a:hover { text-decoration: underline; }

        #stats-badge {
            background: #dbeafe;
            color: #1e40af;
            padding: 0.25rem 0.75rem;
            border-radius: 9999px;
            font-size: 0.875rem;
            font-weight: 600;
        }
    </style>
</head>
<body>
    <div id="header">
        <h1>Tbilisi Education Explorer</h1>
        <button onclick="downloadCSV()">
            <svg width="20" height="20" fill="none" stroke="currentColor" viewBox="0 0 24 24"><path stroke-linecap="round" stroke-linejoin="round" stroke-width="2" d="M4 16v1a3 3 0 003 3h10a3 3 0 003-3v-1m-4-4l-4 4m0 0l-4-4m4 4V4"></path></svg>
            Download CSV
        </button>
    </div>
    <div id="controls">
        <input type="text" id="search" placeholder="Search by name..." onkeyup="filterData()">
        <select id="categoryFilter" onchange="filterData()">
            <option value="all">All Categories</option>
            <option value="School">School</option>
            <option value="University">University</option>
            <option value="College">College</option>
            <option value="Academy/Learning Center">Academy/Center</option>
            <option value="Kindergarten">Kindergarten</option>
        </select>
        <div id="stats-badge">Loading...</div>
    </div>
    <div id="map"></div>

    <script src="https://unpkg.com/leaflet@1.9.4/dist/leaflet.js"></script>
    <script>
        const data = __DATA_JSON__;
        let map = L.map('map', { zoomControl: false }).setView([41.6938, 44.8015], 13);
        
        L.control.zoom({ position: 'bottomleft' }).addTo(map);

        L.tileLayer('https://{s}.basemaps.cartocdn.com/rastertiles/voyager/{z}/{x}/{y}{r}.png', {
            attribution: '&copy; <a href="https://www.openstreetmap.org/copyright">OpenStreetMap</a> contributors &copy; <a href="https://carto.com/attributions">CARTO</a>',
            subdomains: 'abcd',
            maxZoom: 20
        }).addTo(map);

        const markers = [];
        const colors = {
            'School': '#3b82f6',
            'University': '#ef4444',
            'Academy/Learning Center': '#10b981',
            'College': '#f59e0b',
            'Kindergarten': '#8b5cf6',
            'Other': '#64748b'
        };

        function initMarkers() {
            data.forEach(item => {
                const color = colors[item.category] || colors['Other'];
                const marker = L.circleMarker([item.lat, item.lng], {
                    radius: 10,
                    fillColor: color,
                    color: "white",
                    weight: 2,
                    opacity: 1,
                    fillOpacity: 0.9
                }).addTo(map);

                marker.bindPopup(`
                    <div class="popup-card">
                        <h3>${item.name}</h3>
                        <p><strong>Category:</strong> ${item.category}</p>
                        <p><strong>Phone:</strong> ${item.phone}</p>
                        <p><strong>Address:</strong> ${item.address}</p>
                        <p><a href="${item.website}" target="_blank">Visit Website</a> | <a href="${item.google_maps_url}" target="_blank">Maps</a></p>
                    </div>
                `);
                
                marker.category = item.category;
                marker.name = item.name.toLowerCase();
                markers.push(marker);
            });
            updateStats(data.length);
        }

        function filterData() {
            const searchText = document.getElementById('search').value.toLowerCase();
            const catFilter = document.getElementById('categoryFilter').value;
            let count = 0;

            markers.forEach(m => {
                const matchesSearch = m.name.includes(searchText);
                const matchesCat = catFilter === 'all' || m.category === catFilter;

                if (matchesSearch && matchesCat) {
                    m.addTo(map);
                    count++;
                } else {
                    map.removeLayer(m);
                }
            });
            updateStats(count);
        }

        function updateStats(count) {
            document.getElementById('stats-badge').innerText = `Results: ${count}`;
        }

        const legend = L.control({position: 'bottomright'});
        legend.onAdd = function (map) {
            const div = L.DomUtil.create('div', 'legend');
            div.innerHTML += '<h4>Categories</h4>';
            for (const [cat, color] of Object.entries(colors)) {
                if (cat === 'Other') continue;
                div.innerHTML += `<div><i style="background: ${color}; color: ${color}"></i> ${cat}</div>`;
            }
            return div;
        };
        legend.addTo(map);

        function downloadCSV() {
            const searchText = document.getElementById('search').value.toLowerCase();
            const catFilter = document.getElementById('categoryFilter').value;
            
            const filtered = data.filter(item => {
                const matchesSearch = item.name.toLowerCase().includes(searchText);
                const matchesCat = catFilter === 'all' || item.category === catFilter;
                return matchesSearch && matchesCat;
            });

            const headers = ['Name', 'Address', 'Phone', 'Website', 'Rating', 'Category'];
            let csvContent = "\uFEFF"; 
            csvContent += headers.join(",") + "\\n";

            filtered.forEach(item => {
                const row = [
                    `"${item.name}"`,
                    `"${item.address}"`,
                    `"${item.phone}"`,
                    `"${item.website}"`,
                    item.rating,
                    `"${item.category}"`
                ];
                csvContent += row.join(",") + "\\n";
            });

            const blob = new Blob([csvContent], { type: 'text/csv;charset=utf-8;' });
            const link = document.createElement("a");
            link.href = URL.createObjectURL(blob);
            link.download = "tbilisi_education.csv";
            link.click();
        }

        initMarkers();
    </script>
</body>
</html>
"""
    final_html = html_template.replace('__DATA_JSON__', json.dumps(data, ensure_ascii=False))
    with open('output.html', 'w', encoding='utf-8') as f:
        f.write(final_html)

if __name__ == "__main__":
    save_to_excel(demo_data)
    generate_html(demo_data)
    print("Demo files generated.")
