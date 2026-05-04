import os
import requests
import time
import json
import pandas as pd
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule

def log(msg):
    try:
        print(msg, flush=True)
    except:
        pass

def load_env():
    if os.path.exists('.env'):
        with open('.env') as f:
            for line in f:
                if '=' in line:
                    key, value = line.strip().split('=', 1)
                    os.environ[key] = value

load_env()
API_KEY = os.getenv('GOOGLE_API_KEY') or os.getenv('GOOGLE_PLACES_API_KEY')
DISTRICTS = {
    "Vake": (41.7020, 44.7650),
    "Saburtalo": (41.7270, 44.7780),
    "Mtatsminda": (41.6940, 44.7950),
    "Didube": (41.7380, 44.8020),
    "Gldani": (41.7700, 44.8100),
    "Nadzaladevi": (41.7480, 44.8300),
    "Isani": (41.6880, 44.8350),
    "Samgori": (41.6650, 44.8500),
    "Chughureti": (41.7050, 44.8100),
    "Krtsanisi": (41.6700, 44.8150),
    "Didi Dighomi": (41.7600, 44.7500),
    "Ortachala": (41.6780, 44.8250)
}
RADIUS = 3000.0
# New API Types for Construction/Architecture
TYPES = ["architect", "general_contractor", "real_estate_agency"]
SEARCH_QUERIES = [
    "architecture firms Tbilisi", "construction companies Tbilisi", "interior design Tbilisi",
    "სამშენებლო კომპანია თბილისი", "არქიტექტურული სტუდია", "რემონტი თბილისი",
    "building renovation Tbilisi", "engineering company Tbilisi",
    "Metrix Tbilisi", "Renox Tbilisi", "Anagi construction", "Archi construction",
    "m2 development Tbilisi", "Simetria construction", "Domus architecture",
    "სამშენებლო კომპანიები", "არქიტექტურული ბიურო", "ბინების რემონტი",
    "interior design studio Tbilisi", "landscape architecture Tbilisi",
    "real estate development Tbilisi", "structural engineering Tbilisi"
]

HEADERS = {
    'Content-Type': 'application/json',
    'X-Goog-Api-Key': API_KEY,
    'X-Goog-FieldMask': 'places.id,places.displayName,places.formattedAddress,places.internationalPhoneNumber,places.websiteUri,places.googleMapsUri,places.location,places.rating,places.userRatingCount,places.types,places.businessStatus'
}

def fetch_all_data():
    all_data = {}
    place_ids = set()

    # Search Nearby (Grid)
    for district, (lat, lng) in DISTRICTS.items():
        for p_type in TYPES:
            log(f"Searching type [{p_type}] in [{district}]...")
            url = 'https://places.googleapis.com/v1/places:searchNearby'
            data = {
                "includedTypes": [p_type],
                "locationRestriction": {
                    "circle": {
                        "center": {"latitude": lat, "longitude": lng},
                        "radius": RADIUS
                    }
                }
            }
            try:
                res = requests.post(url, headers=HEADERS, json=data).json()
                process_results(res.get('places', []), all_data, place_ids, district)
            except Exception as e:
                log(f"Error searching nearby: {e}")
            time.sleep(0.5)

    # Search Text (Queries)
    for query in SEARCH_QUERIES:
        log(f"Searching query [{query}]...")
        url = 'https://places.googleapis.com/v1/places:searchText'
        data = {"textQuery": query}
        try:
            res = requests.post(url, headers=HEADERS, json=data).json()
            process_results(res.get('places', []), all_data, place_ids, "Unknown")
        except Exception as e:
            log(f"Error searching text: {e}")
            
        time.sleep(0.5)

    log(f"Found {len(all_data)} unique records. Processing data...")
    
    final_data = []
    for p_id, res in all_data.items():
        name = res.get('displayName', {}).get('text', 'N/A')
        address = res.get('formattedAddress', 'N/A')
        
        cleaned = {
            "place_id": p_id,
            "name": name,
            "address": address,
            "phone": res.get("internationalPhoneNumber", "N/A"),
            "website": res.get("websiteUri", "N/A"),
            "google_maps_url": res.get("googleMapsUri", "N/A"),
            "lat": res.get("location", {}).get("latitude"),
            "lng": res.get("location", {}).get("longitude"),
            "rating": res.get("rating", 0),
            "user_ratings_total": res.get("userRatingCount", 0),
            "business_status": res.get("businessStatus", "UNKNOWN"),
            "district": res.get("district", "Unknown")
        }
        
        # Categorization logic
        name_lower = name.lower()
        types = res.get('types', [])
        
        if 'architect' in types or any(k in name_lower for k in ['architecture', 'არქიტექტურა', 'architectural']):
            cleaned['category'] = 'Architecture'
        elif any(k in name_lower for k in ['interior', 'design', 'დიზაინი']):
            cleaned['category'] = 'Interior Design'
        elif 'general_contractor' in types or any(k in name_lower for k in ['construction', 'building', 'მშენებლობა', 'development']):
            cleaned['category'] = 'Construction'
        elif any(k in name_lower for k in ['renovation', 'repair', 'რემონტი']):
            cleaned['category'] = 'Renovation'
        else:
            cleaned['category'] = 'Engineering/Other'
        
        final_data.append(cleaned)

    # Save to JSON
    with open('tbilisi_construction_clean.json', 'w', encoding='utf-8') as f:
        json.dump(final_data, f, ensure_ascii=False, indent=2)

    return final_data

def process_results(places, all_places, place_ids, district):
    for p in places:
        p_id = p['id']
        if p_id not in place_ids:
            place_ids.add(p_id)
            p['district'] = district
            all_places[p_id] = p

def save_to_excel(data):
    if not data: return
    # Filtered export data
    export_data = []
    for item in data:
        export_data.append({
            "Name": item["name"],
            "Category": item["category"],
            "District": item["district"],
            "Address": item["address"],
            "Phone": item["phone"],
            "Website": item["website"],
            "Google Maps Link": item["google_maps_url"],
            "Rating": item["rating"],
            "Votes (Number of Ratings)": item["user_ratings_total"],
            "Business Status": item["business_status"]
        })
        
    df = pd.DataFrame(export_data)
    writer = pd.ExcelWriter('construction_architecture.xlsx', engine='openpyxl')
    
    df.to_excel(writer, sheet_name='All Records', index=False)
    ws = writer.sheets['All Records']
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    
    # Styling
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1e293b", end_color="1e293b", fill_type="solid")

    for cat in df['Category'].unique():
        cat_df = df[df['Category'] == cat]
        sheet_name = str(cat).replace('/', '-')[:30]
        cat_df.to_excel(writer, sheet_name=sheet_name, index=False)
    
    writer.close()

TEMPLATE = r"""
<!DOCTYPE html>
<html lang="ka">
<head>
    <meta charset="UTF-8">
    <title>Construction & Architecture Explorer</title>
    <link rel="stylesheet" href="https://unpkg.com/leaflet@1.9.4/dist/leaflet.css" />
    <link rel="stylesheet" href="https://unpkg.com/leaflet.markercluster@1.4.1/dist/MarkerCluster.css" />
    <link rel="stylesheet" href="https://unpkg.com/leaflet.markercluster@1.4.1/dist/MarkerCluster.Default.css" />
    <link href="https://fonts.googleapis.com/css2?family=Outfit:wght@300;400;600&display=swap" rel="stylesheet">
    <style>
        :root { --primary: #0f172a; --secondary: #334155; --bg: #f8fafc; }
        body { font-family: 'Outfit', sans-serif; margin: 0; display: flex; flex-direction: column; height: 100vh; background: var(--bg); }
        #header { background: var(--primary); color: white; padding: 0.8rem 1.5rem; display: flex; justify-content: space-between; align-items: center; z-index: 1001; }
        #main-container { display: flex; flex: 1; overflow: hidden; }
        #sidebar { width: 350px; background: white; border-right: 1px solid #e2e8f0; display: flex; flex-direction: column; }
        #controls { padding: 1rem; border-bottom: 1px solid #e2e8f0; background: #fff; }
        #list-view { flex: 1; overflow-y: auto; padding: 0.5rem; }
        #map { flex: 1; }
        .location-item { padding: 0.8rem; border-bottom: 1px solid #f1f5f9; cursor: pointer; transition: 0.2s; }
        .location-item:hover { background: #f8fafc; }
        .location-item h4 { margin: 0 0 0.3rem 0; font-size: 0.95rem; }
        .location-item p { margin: 0; font-size: 0.8rem; color: #64748b; }
        .badge { display: inline-block; padding: 2px 8px; border-radius: 4px; font-size: 0.7rem; font-weight: 600; color: white; margin-top: 4px; }
        button { padding: 0.6rem 1rem; border-radius: 8px; border: none; background: #334155; color: white; font-weight: 600; cursor: pointer; font-size: 0.8rem; }
        select, input { width: 100%; padding: 0.5rem; border: 1px solid #e2e8f0; border-radius: 6px; font-size: 0.85rem; margin-bottom: 0.5rem; }
        #stats-bar { padding: 0.5rem 1rem; background: #e2e8f0; font-size: 0.8rem; font-weight: 600; color: #1e293b; }

        /* Premium Popup Styles */
        .leaflet-popup-content-wrapper { padding: 0; overflow: hidden; border-radius: 12px; }
        .leaflet-popup-content { margin: 0; width: 250px !important; }
        .popup-card { display: flex; flex-direction: column; }
        .popup-header { padding: 12px; color: white; }
        .popup-header h3 { margin: 0; font-size: 1rem; font-weight: 600; }
        .popup-body { padding: 12px; font-size: 0.85rem; color: #334155; line-height: 1.4; }
        .popup-body p { margin: 4px 0; }
        .popup-footer { display: flex; gap: 8px; padding: 12px; background: #f8fafc; border-top: 1px solid #e2e8f0; }
        .popup-btn { flex: 1; padding: 6px; border-radius: 6px; text-decoration: none; font-size: 0.75rem; font-weight: 600; text-align: center; transition: 0.2s; }
        .popup-btn.maps { background: #334155; color: white; }
        .popup-btn.web { background: #2563eb; color: white; }
        .popup-btn.disabled { background: #e2e8f0; color: #94a3b8; cursor: not-allowed; }
        .popup-btn:hover:not(.disabled) { opacity: 0.9; transform: translateY(-1px); }
    </style>
</head>
<body>
    <div id="header">
        <h1 style="margin:0; font-size:1.2rem;">Construction & Architecture Explorer</h1>
        <button onclick="downloadXLSX()">Download XLSX</button>
    </div>
    <div id="stats-bar">Loading data...</div>
    <div id="main-container">
        <div id="sidebar">
            <div id="controls">
                <select id="catFilter" onchange="applyFilters()">
                    <option value="all">All Categories</option>
                    <option value="Architecture">Architecture</option>
                    <option value="Construction">Construction</option>
                    <option value="Interior Design">Interior Design</option>
                    <option value="Renovation">Renovation</option>
                </select>
                <input type="text" id="searchInput" placeholder="Search by name..." onkeyup="applyFilters()">
            </div>
            <div id="list-view"></div>
        </div>
        <div id="map"></div>
    </div>

    <script src="https://unpkg.com/leaflet@1.9.4/dist/leaflet.js"></script>
    <script src="https://unpkg.com/leaflet.markercluster@1.4.1/dist/leaflet.markercluster.js"></script>
    <script src="https://cdn.sheetjs.com/xlsx-0.20.1/package/dist/xlsx.full.min.js"></script>
    <script>
        const data = __DATA_JSON__;
        const colors = {
            'Architecture': '#2563eb', 'Construction': '#ef4444', 
            'Interior Design': '#db2777', 'Renovation': '#059669', 'Engineering/Other': '#64748b'
        };

        const map = L.map('map').setView([41.6938, 44.8015], 13);
        L.tileLayer('https://{s}.basemaps.cartocdn.com/rastertiles/voyager/{z}/{x}/{y}{r}.png').addTo(map);
        
        const cluster = L.markerClusterGroup({
            spiderfyOnMaxZoom: true,
            showCoverageOnHover: false,
            zoomToBoundsOnClick: true,
            maxClusterRadius: 40
        });
        map.addLayer(cluster);

        let filteredData = [...data];
        const markersMap = new Map();

        function init() {
            renderList();
            renderMarkers();
        }

        function renderMarkers() {
            cluster.clearLayers();
            filteredData.forEach(item => {
                const color = colors[item.category] || '#64748b';
                const marker = L.circleMarker([item.lat, item.lng], {
                    radius: 8, fillColor: color, color: '#fff', weight: 1, opacity: 1, fillOpacity: 0.9
                });

                const hasWeb = item.website && item.website !== 'N/A';
                
                marker.bindPopup(`
                    <div class="popup-card">
                        <div class="popup-header" style="background: ${color}">
                            <h3>${item.name}</h3>
                        </div>
                        <div class="popup-body">
                            <p><b>Category:</b> ${item.category}</p>
                            <p><b>Address:</b> ${item.address}</p>
                            <p><b>Phone:</b> ${item.phone}</p>
                            ${item.rating ? `<p><b>Rating:</b> ⭐${item.rating} (${item.user_ratings_total})</p>` : ''}
                        </div>
                        <div class="popup-footer">
                            <a href="${item.google_maps_url}" target="_blank" class="popup-btn maps">Maps</a>
                            ${hasWeb ? `<a href="${item.website}" target="_blank" class="popup-btn web">Website</a>` : '<span class="popup-btn disabled">No Web</span>'}
                        </div>
                    </div>
                `);
                cluster.addLayer(marker);
                markersMap.set(item.place_id, marker);
            });
            updateStats();
        }

        function renderList() {
            const list = document.getElementById('list-view');
            list.innerHTML = filteredData.map(item => `
                <div class="location-item" onclick="focusLocation('${item.place_id}')">
                    <h4>${item.name}</h4>
                    <p>${item.address}</p>
                    <span class="badge" style="background:${colors[item.category] || '#64748b'}">${item.category}</span>
                </div>
            `).join('');
        }

        function applyFilters() {
            const cat = document.getElementById('catFilter').value;
            const search = document.getElementById('searchInput').value.toLowerCase();

            filteredData = data.filter(item => {
                const matchesCat = cat === 'all' || item.category === cat;
                const matchesSearch = item.name.toLowerCase().includes(search);
                return matchesCat && matchesSearch;
            });

            renderList();
            renderMarkers();
        }

        function focusLocation(id) {
            const marker = markersMap.get(id);
            if (marker) {
                map.setView(marker.getLatLng(), 16);
                marker.openPopup();
            }
        }

        function updateStats() {
            document.getElementById('stats-bar').innerText = `Total Records Found: ${data.length} | Currently Visible: ${filteredData.length}`;
        }

        function downloadXLSX() {
            const exportData = filteredData.map(i => ({
                'Name': i.name,
                'Category': i.category,
                'District': i.district,
                'Address': i.address,
                'Phone': i.phone,
                'Website': i.website,
                'Google Maps Link': i.google_maps_url,
                'Rating': i.rating
            }));

            const ws = XLSX.utils.json_to_sheet(exportData);
            const wb = XLSX.utils.book_new();
            XLSX.utils.book_append_sheet(wb, ws, "Construction");
            XLSX.writeFile(wb, "tbilisi_construction_architecture.xlsx");
        }

        init();
    </script>
</body>
</html>
"""

def generate_html(data):
    final_html = TEMPLATE.replace('__DATA_JSON__', json.dumps(data, ensure_ascii=False))
    with open('construction_explorer.html', 'w', encoding='utf-8') as f:
        f.write(final_html)

if __name__ == "__main__":
    if not API_KEY:
        log("Error: GOOGLE_PLACES_API_KEY not set.")
    else:
        results = fetch_all_data()
        save_to_excel(results)
        generate_html(results)
        log("Success! Created construction_architecture.xlsx and construction_explorer.html")
