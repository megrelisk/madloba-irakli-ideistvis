import os
import requests
import time
import json
import pandas as pd
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule

def log(msg):
    try:
        print(msg, flush=True)
    except:
        pass

def load_env():
    if os.path.exists('.env'):
        with open('.env') as f:
            for line in f:
                if '=' in line:
                    key, value = line.strip().split('=', 1)
                    os.environ[key] = value

load_env()
API_KEY = os.getenv('GOOGLE_API_KEY') or os.getenv('GOOGLE_PLACES_API_KEY')
DISTRICTS = {
    "Vake": (41.7020, 44.7650),
    "Saburtalo": (41.7270, 44.7780),
    "Mtatsminda": (41.6940, 44.7950),
    "Didube": (41.7380, 44.8020),
    "Gldani": (41.7700, 44.8100),
    "Nadzaladevi": (41.7480, 44.8300),
    "Isani": (41.6880, 44.8350),
    "Samgori": (41.6650, 44.8500),
    "Chughureti": (41.7050, 44.8100),
    "Krtsanisi": (41.6700, 44.8150),
    "Didi Dighomi": (41.7600, 44.7500),
    "Ortachala": (41.6780, 44.8250)
}
RADIUS = 3000.0
# New API Types: https://developers.google.com/maps/documentation/places/web-service/place-types
TYPES = ["school", "university", "library", "preschool", "secondary_school"]
SEARCH_QUERIES = [
    "სკოლა თბილისში", "უნივერსიტეტი თბილისში", "კოლეჯი თბილისში", "აკადემია თბილისში", 
    "საბავშვო ბაღი თბილისი", "სასწავლო ცენტრი თბილისი", "tutoring center Tbilisi",
    "language school Tbilisi", "music school Tbilisi", "art school Tbilisi"
]

HEADERS = {
    'Content-Type': 'application/json',
    'X-Goog-Api-Key': API_KEY,
    'X-Goog-FieldMask': 'places.id,places.displayName,places.formattedAddress,places.internationalPhoneNumber,places.websiteUri,places.googleMapsUri,places.location,places.rating,places.userRatingCount,places.regularOpeningHours,places.types,places.businessStatus'
}

def fetch_all_data():
    all_places = {}
    place_ids = set()

    # Search Nearby with Types (Grid)
    for district, (lat, lng) in DISTRICTS.items():
        for p_type in TYPES:
            log(f"Searching type [{p_type}] in [{district}]...")
            url = 'https://places.googleapis.com/v1/places:searchNearby'
            data = {
                "includedTypes": [p_type],
                "locationRestriction": {
                    "circle": {
                        "center": {"latitude": lat, "longitude": lng},
                        "radius": RADIUS
                    }
                }
            }
            res = requests.post(url, headers=HEADERS, json=data).json()
            process_results(res.get('places', []), all_places, place_ids, district)
            time.sleep(0.5)

    # Search Text (Queries)
    for query in SEARCH_QUERIES:
        log(f"Searching query [{query}]...")
        url = 'https://places.googleapis.com/v1/places:searchText'
        data = {"textQuery": query}
        res = requests.post(url, headers=HEADERS, json=data).json()
        process_results(res.get('places', []), all_places, place_ids, "Unknown")
        time.sleep(0.5)

    log(f"Found {len(all_places)} unique places. Processing data...")
    
    final_data = []
    for p_id, res in all_places.items():
        name = res.get('displayName', {}).get('text', 'N/A')
        address = res.get('formattedAddress', 'N/A')
        
        cleaned = {
            "place_id": p_id,
            "name": name,
            "address": address,
            "phone": res.get("internationalPhoneNumber", "N/A"),
            "website": res.get("websiteUri", "N/A"),
            "google_maps_url": res.get("googleMapsUri", "N/A"),
            "lat": res.get("location", {}).get("latitude"),
            "lng": res.get("location", {}).get("longitude"),
            "rating": res.get("rating", 0),
            "user_ratings_total": res.get("userRatingCount", 0),
            "business_status": res.get("businessStatus", "UNKNOWN"),
            "district": res.get("district", "Unknown"),
            "opening_hours": ", ".join(res.get('regularOpeningHours', {}).get('weekdayDescriptions', [])) if res.get('regularOpeningHours') else 'N/A'
        }
        
        # Categorization
        name_lower = name.lower()
        types = res.get('types', [])
        if 'university' in types or 'უნივერსიტეტი' in name_lower or 'university' in name_lower:
            cleaned['category'] = 'University'
        elif 'kindergarten' in types or any(k in name_lower for k in ['ბაღი', 'kindergarten', 'preschool', 'საბავშვო']):
            cleaned['category'] = 'Kindergarten'
        elif 'academy' in name_lower or 'აკადემია' in name_lower:
            cleaned['category'] = 'Academy'
        elif any(k in name_lower for k in ['ცენტრი', 'center', 'სასწავლო', 'tutoring']):
            cleaned['category'] = 'Learning Center'
        elif 'school' in types or 'სკოლა' in name_lower or 'school' in name_lower:
            cleaned['category'] = 'School'
        else:
            cleaned['category'] = 'Other'
        
        final_data.append(cleaned)

    # Save to JSON
    with open('tbilisi_schools_clean.json', 'w', encoding='utf-8') as f:
        json.dump(final_data, f, ensure_ascii=False, indent=2)

    return final_data

def process_results(places, all_places, place_ids, district):
    for p in places:
        p_id = p['id']
        if p_id not in place_ids:
            place_ids.add(p_id)
            p['district'] = district
            all_places[p_id] = p

def save_to_excel(data):
    if not data: return
    df = pd.DataFrame(data)
    writer = pd.ExcelWriter('tbilisi_education_full.xlsx', engine='openpyxl')
    
    df.to_excel(writer, sheet_name='All Locations', index=False)
    ws = writer.sheets['All Locations']
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    
    # Conditional formatting
    green_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    max_r = ws.max_row
    if max_r > 1:
        ws.conditional_formatting.add(f'I2:I{max_r}', CellIsRule(operator='greaterThanOrEqual', formula=['4.5'], fill=green_fill))

    for district in df['district'].unique():
        dist_df = df[df['district'] == district]
        dist_df.to_excel(writer, sheet_name=str(district)[:30], index=False)
    
    writer.close()

def generate_html(data):
    # (Same HTML generation logic as before, just using the new field names if needed)
    # The fields in final_data are the same as before.
    import json
    with open('output.html', 'w', encoding='utf-8') as f:
        # Re-using the same template logic
        dist_options = "\n".join([f'<option value="{d}">{d}</option>' for d in DISTRICTS.keys()])
        final_html = TEMPLATE.replace('__DATA_JSON__', json.dumps(data, ensure_ascii=False))
        final_html = final_html.replace('__DISTRICT_OPTIONS__', dist_options)
        f.write(final_html)

TEMPLATE = """
<!DOCTYPE html>
<html lang="ka">
<head>
    <meta charset="UTF-8">
    <title>Tbilisi Education Explorer PRO</title>
    <link rel="stylesheet" href="https://unpkg.com/leaflet@1.9.4/dist/leaflet.css" />
    <link rel="stylesheet" href="https://unpkg.com/leaflet.markercluster@1.4.1/dist/MarkerCluster.css" />
    <link rel="stylesheet" href="https://unpkg.com/leaflet.markercluster@1.4.1/dist/MarkerCluster.Default.css" />
    <link href="https://fonts.googleapis.com/css2?family=Outfit:wght@300;400;600&display=swap" rel="stylesheet">
    <style>
        :root { --primary: #2563eb; --secondary: #1e293b; --bg: #f8fafc; }
        body { font-family: 'Outfit', sans-serif; margin: 0; display: flex; flex-direction: column; height: 100vh; background: var(--bg); }
        #header { background: var(--secondary); color: white; padding: 0.8rem 1.5rem; display: flex; justify-content: space-between; align-items: center; z-index: 1001; }
        #main-container { display: flex; flex: 1; overflow: hidden; }
        #sidebar { width: 350px; background: white; border-right: 1px solid #e2e8f0; display: flex; flex-direction: column; }
        #controls { padding: 1rem; border-bottom: 1px solid #e2e8f0; background: #fff; }
        #list-view { flex: 1; overflow-y: auto; padding: 0.5rem; }
        #map { flex: 1; }
        .location-item { padding: 0.8rem; border-bottom: 1px solid #f1f5f9; cursor: pointer; transition: 0.2s; }
        .location-item:hover { background: #f8fafc; }
        .location-item h4 { margin: 0 0 0.3rem 0; font-size: 0.95rem; }
        .location-item p { margin: 0; font-size: 0.8rem; color: #64748b; }
        .badge { display: inline-block; padding: 2px 8px; border-radius: 4px; font-size: 0.7rem; font-weight: 600; color: white; margin-top: 4px; }
        .filter-group { margin-bottom: 1rem; }
        .filter-group label { display: block; font-size: 0.8rem; font-weight: 600; color: #64748b; margin-bottom: 0.4rem; }
        select, input { width: 100%; padding: 0.5rem; border: 1px solid #e2e8f0; border-radius: 6px; font-size: 0.85rem; margin-bottom: 0.5rem; }
        .cat-checks { display: grid; grid-template-columns: 1fr 1fr; gap: 0.3rem; }
        .cat-checks label { font-size: 0.75rem; font-weight: normal; cursor: pointer; }
        .btn-group { display: flex; gap: 0.5rem; }
        button { flex: 1; padding: 0.5rem; border-radius: 6px; border: none; background: var(--primary); color: white; font-weight: 600; cursor: pointer; font-size: 0.8rem; }
        button.secondary { background: #64748b; }
        #stats-bar { padding: 0.5rem 1rem; background: #dbeafe; font-size: 0.8rem; font-weight: 600; color: #1e40af; }
        
        #preview-table-container { position: absolute; bottom: 0; left: 350px; right: 0; background: white; max-height: 200px; overflow-y: auto; z-index: 1000; border-top: 2px solid var(--primary); display: none; }
        table { width: 100%; border-collapse: collapse; font-size: 0.75rem; }
        th, td { padding: 8px; text-align: left; border-bottom: 1px solid #eee; }
        th { background: #f8fafc; }
    </style>
</head>
<body>
    <div id="header">
        <h1 style="margin:0; font-size:1.2rem;">Tbilisi Education Explorer PRO</h1>
        <div class="btn-group" style="width: auto;">
            <button onclick="downloadCSV('all')">Download All CSV</button>
            <button onclick="downloadCSV('filtered')" class="secondary">Download Filtered</button>
            <button onclick="togglePreview()" class="secondary">Preview Table</button>
        </div>
    </div>
    <div id="stats-bar">Loading locations...</div>
    <div id="main-container">
        <div id="sidebar">
            <div id="controls">
                <div class="filter-group">
                    <label>District</label>
                    <select id="distFilter" onchange="applyFilters()">
                        <option value="all">All Districts</option>
                        __DISTRICT_OPTIONS__
                    </select>
                </div>
                <div class="filter-group">
                    <label>Categories</label>
                    <div class="cat-checks">
                        <label><input type="checkbox" value="School" checked onchange="applyFilters()"> School</label>
                        <label><input type="checkbox" value="University" checked onchange="applyFilters()"> University</label>
                        <label><input type="checkbox" value="Kindergarten" checked onchange="applyFilters()"> Kindergarten</label>
                        <label><input type="checkbox" value="Academy" checked onchange="applyFilters()"> Academy</label>
                        <label><input type="checkbox" value="Learning Center" checked onchange="applyFilters()"> Center</label>
                        <label><input type="checkbox" value="Other" checked onchange="applyFilters()"> Other</label>
                    </div>
                </div>
                <input type="text" id="searchInput" placeholder="Search by name..." onkeyup="applyFilters()">
            </div>
            <div id="list-view"></div>
        </div>
        <div id="map"></div>
    </div>

    <div id="preview-table-container">
        <table id="preview-table">
            <thead>
                <tr><th>Name</th><th>Category</th><th>District</th><th>Rating</th><th>Phone</th></tr>
            </thead>
            <tbody id="preview-body"></tbody>
        </table>
    </div>

    <script src="https://unpkg.com/leaflet@1.9.4/dist/leaflet.js"></script>
    <script src="https://unpkg.com/leaflet.markercluster@1.4.1/dist/leaflet.markercluster.js"></script>
    <script>
        const data = __DATA_JSON__;
        const colors = {
            'School': '#3b82f6', 'University': '#ef4444', 'Kindergarten': '#f59e0b',
            'Academy': '#8b5cf6', 'Learning Center': '#10b981', 'Other': '#64748b'
        };

        const map = L.map('map').setView([41.6938, 44.8015], 12);
        L.tileLayer('https://{s}.basemaps.cartocdn.com/rastertiles/voyager/{z}/{x}/{y}{r}.png').addTo(map);
        
        const cluster = L.markerClusterGroup();
        map.addLayer(cluster);

        let filteredData = [...data];
        const markersMap = new Map();

        function init() {
            renderList();
            renderMarkers();
            renderPreview();
        }

        function renderMarkers() {
            cluster.clearLayers();
            filteredData.forEach(item => {
                const color = colors[item.category] || colors['Other'];
                const marker = L.circleMarker([item.lat, item.lng], {
                    radius: 8, fillColor: color, color: '#fff', weight: 1, opacity: 1, fillOpacity: 0.9
                });
                marker.bindPopup(`
                    <div style="min-width:180px">
                        <h4 style="margin:0 0 5px 0">${item.name}</h4>
                        <p style="font-size:0.8rem; margin:2px 0"><b>Category:</b> ${item.category}</p>
                        <p style="font-size:0.8rem; margin:2px 0"><b>District:</b> ${item.district}</p>
                        <p style="font-size:0.8rem; margin:2px 0"><b>Rating:</b> ⭐${item.rating} (${item.user_ratings_total})</p>
                        <p style="font-size:0.8rem; margin:2px 0"><b>Phone:</b> ${item.phone}</p>
                        <a href="${item.website}" target="_blank" style="font-size:0.8rem">Website</a> | 
                        <a href="${item.google_maps_url}" target="_blank" style="font-size:0.8rem">Maps</a>
                    </div>
                `);
                cluster.addLayer(marker);
                markersMap.set(item.place_id, marker);
            });
            updateStats();
        }

        function renderList() {
            const list = document.getElementById('list-view');
            list.innerHTML = filteredData.map(item => `
                <div class="location-item" onclick="focusLocation('${item.place_id}')">
                    <h4>${item.name}</h4>
                    <p>${item.address}</p>
                    <span class="badge" style="background:${colors[item.category] || colors['Other']}">${item.category}</span>
                </div>
            `).join('');
        }

        function renderPreview() {
            const body = document.getElementById('preview-body');
            body.innerHTML = filteredData.slice(0, 15).map(item => `
                <tr>
                    <td>${item.name}</td>
                    <td>${item.category}</td>
                    <td>${item.district}</td>
                    <td>⭐${item.rating}</td>
                    <td>${item.phone}</td>
                </tr>
            `).join('');
        }

        function applyFilters() {
            const dist = document.getElementById('distFilter').value;
            const search = document.getElementById('searchInput').value.toLowerCase();
            const checkedCats = Array.from(document.querySelectorAll('.cat-checks input:checked')).map(i => i.value);

            filteredData = data.filter(item => {
                const matchesDist = dist === 'all' || item.district === dist;
                const matchesCat = checkedCats.includes(item.category);
                const matchesSearch = item.name.toLowerCase().includes(search);
                return matchesDist && matchesCat && matchesSearch;
            });

            renderList();
            renderMarkers();
            renderPreview();
        }

        function focusLocation(id) {
            const marker = markersMap.get(id);
            if (marker) {
                map.setView(marker.getLatLng(), 16);
                marker.openPopup();
            }
        }

        function updateStats() {
            document.getElementById('stats-bar').innerText = `Total Locations Found: ${data.length} | Currently Visible: ${filteredData.length}`;
        }

        function togglePreview() {
            const p = document.getElementById('preview-table-container');
            p.style.display = p.style.display === 'none' ? 'block' : 'none';
        }

        function downloadCSV(type) {
            const items = type === 'all' ? data : filteredData;
            const headers = ['Name', 'Category', 'District', 'Address', 'Phone', 'Website', 'Rating', 'Reviews'];
            let csv = "\\uFEFF" + headers.join(",") + "\\n";
            items.forEach(i => {
                const row = [`"${i.name}"`, `"${i.category}"`, `"${i.district}"`, `"${i.address}"`, `"${i.phone}"`, `"${i.website}"`, i.rating, i.user_ratings_total];
                csv += row.join(",") + "\\n";
            });
            const blob = new Blob([csv], { type: 'text/csv;charset=utf-8;' });
            const link = document.createElement("a");
            link.href = URL.createObjectURL(blob);
            link.download = `tbilisi_edu_${type}.csv`;
            link.click();
        }

        init();
    </script>
</body>
</html>
"""

def generate_html(data):
    dist_options = "\n".join([f'<option value="{d}">{d}</option>' for d in DISTRICTS.keys()])
    final_html = TEMPLATE.replace('__DATA_JSON__', json.dumps(data, ensure_ascii=False))
    final_html = final_html.replace('__DISTRICT_OPTIONS__', dist_options)
    with open('output.html', 'w', encoding='utf-8') as f:
        f.write(final_html)

if __name__ == "__main__":
    if not API_KEY:
        log("Error: GOOGLE_PLACES_API_KEY not set.")
    else:
        results = fetch_all_data()
        save_to_excel(results)
        generate_html(results)
        log("Success! Created tbilisi_education_full.xlsx and output.html")
